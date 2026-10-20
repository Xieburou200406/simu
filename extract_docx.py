import sys, io
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

def extract(path):
    doc = Document(path)
    out = []
    # 段落
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt:
            style = p.style.name if p.style else ''
            out.append(f"[{style}] {txt}")
    # 表格
    for ti, table in enumerate(doc.tables):
        out.append(f"\n===== TABLE {ti+1} =====")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            out.append(" | ".join(cells))
    return "\n".join(out)

if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2]
    text = extract(src)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"WROTE {dst}  chars={len(text)}  lines={text.count(chr(10))+1}")
